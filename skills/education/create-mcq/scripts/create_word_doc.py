#!/usr/bin/env python3
"""
create_word_doc.py — Generate a Word (.docx) file from MCQ JSON data.

Usage:
    python create_word_doc.py <input.json> <output.docx>

Input JSON format:
    {
      "title": "MCQ Set",          # optional
      "questions": [
        {
          "number": 1,
          "stem": "Question text",
          "correct_answer": "A",
          "answers": [
            {"letter": "A", "text": "Answer text", "explanation": "Why correct/incorrect"},
            ...
          ]
        }
      ]
    }
"""

import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Error: python-docx is not installed. Run: pip install python-docx")
    sys.exit(1)


FONT_NAME = "Aptos"
BODY_SIZE = Pt(11)
HEADING_SIZE = Pt(13)
GRAY_HEADER = RGBColor(0xD9, 0xD9, 0xD9)  # Light gray for table header


def set_font(run, size=None, bold=False):
    run.font.name = FONT_NAME
    run.font.bold = bold
    if size:
        run.font.size = size
    # Ensure the font applies to non-Latin characters too
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rFonts.set(qn("w:cs"), FONT_NAME)
    rPr.insert(0, rFonts)


def shade_cell(cell, fill_color):
    """Apply a background fill color to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    hex_color = "{:02X}{:02X}{:02X}".format(fill_color.red, fill_color.green, fill_color.blue)
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_question(doc, question):
    number = question.get("number", "")
    stem = question.get("stem", "")
    correct = question.get("correct_answer", "")
    answers = question.get("answers", [])

    # Question heading
    heading = doc.add_heading(f"Question {number}", level=2)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.font.name = FONT_NAME
        run.font.size = HEADING_SIZE
        run.font.color.rgb = RGBColor(0, 0, 0)

    # Stem
    stem_para = doc.add_paragraph(stem)
    for run in stem_para.runs:
        set_font(run, size=BODY_SIZE)

    # Correct answer line (bold)
    correct_para = doc.add_paragraph()
    run = correct_para.add_run(f"Correct answer: {correct}")
    set_font(run, size=BODY_SIZE, bold=True)

    # Answer table
    if answers:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        # Header row
        hdr_cells = table.rows[0].cells
        headers = ["Letter", "Answer", "Explanation"]
        for i, (cell, header) in enumerate(zip(hdr_cells, headers)):
            shade_cell(cell, GRAY_HEADER)
            para = cell.paragraphs[0]
            run = para.add_run(header)
            set_font(run, size=BODY_SIZE, bold=True)

        # Answer rows
        for answer in answers:
            row_cells = table.add_row().cells
            values = [
                answer.get("letter", ""),
                answer.get("text", ""),
                answer.get("explanation", ""),
            ]
            for cell, value in zip(row_cells, values):
                para = cell.paragraphs[0]
                run = para.add_run(value)
                set_font(run, size=BODY_SIZE)

    # Spacing after each question
    doc.add_paragraph()


def build_document(data, output_path):
    doc = Document()

    # Set default font for the document's normal style
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BODY_SIZE

    # Optional title
    title_text = data.get("title", "").strip()
    if title_text:
        title = doc.add_heading(title_text, level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in title.runs:
            run.font.name = FONT_NAME
            run.font.color.rgb = RGBColor(0, 0, 0)
        doc.add_paragraph()

    for question in data.get("questions", []):
        add_question(doc, question)

    doc.save(output_path)
    print(f"Saved: {output_path}")


def main():
    if len(sys.argv) != 3:
        print("Usage: python create_word_doc.py <input.json> <output.docx>")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    output_path = sys.argv[2]

    if not input_path.exists():
        print(f"Error: Input file not found: {input_path}")
        sys.exit(1)

    with open(input_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    build_document(data, output_path)


if __name__ == "__main__":
    main()
